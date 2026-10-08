"""Builds PayPilot_pitch.pptx (7 slides). Run: python build_deck.py"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

NAVY, TEAL, WHITE, GREY = RGBColor(0x0B, 0x1F, 0x3A), RGBColor(0x1F, 0xC8, 0xB4), RGBColor(255, 255, 255), RGBColor(0xB8, 0xC4, 0xD6)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


def slide(title, bullets=None, sub=None, big=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    bar = s.shapes.add_shape(1, 0, 0, Inches(0.25), prs.slide_height)
    bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.8), Inches(0.5 if not big else 2.4), Inches(11.7), Inches(1.4))
    p = tb.text_frame.paragraphs[0]
    p.text = title; p.font.size = Pt(54 if big else 38); p.font.bold = True; p.font.color.rgb = WHITE
    if sub:
        t2 = s.shapes.add_textbox(Inches(0.8), Inches(3.6 if big else 1.5), Inches(11.7), Inches(1))
        q = t2.text_frame.paragraphs[0]
        q.text = sub; q.font.size = Pt(24); q.font.color.rgb = TEAL
    if bullets:
        b = s.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.7), Inches(4.8))
        tf = b.text_frame; tf.word_wrap = True
        for i, text in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = "•  " + text; para.font.size = Pt(26); para.font.color.rgb = WHITE
            para.space_after = Pt(14)
    return s


slide("PayPilot", sub="The AI agent that chases unpaid invoices - and never sends without your OK",
      big=True)
slide("The problem", [
    "Small UAE businesses lose time and cash chasing late invoices",
    "Chasing is repetitive, awkward and easy to forget",
    "Invoices with mistakes (wrong VAT, missing TRN) make it worse",
    "Owners don't want a robot emailing customers on its own"],
      sub="Late payments hurt cash flow")
slide("What PayPilot does", [
    "1. Reads invoices and checks them (5% VAT, TRN, totals)",
    "2. Finds overdue and soon-due invoices",
    "3. Writes reminders: friendly first, firmer later",
    "4. Manager approves - only then it sends",
    "5. Reads customer replies: promise, dispute, 'already paid'",
    "6. Audit log + weekly cash summary"], sub="Design. Deploy. Delegate.")
slide("Safe by design", [
    "Human approval before any message goes out",
    "The block is in the code, not just the prompt: send is refused until approved",
    "Invoice errors, disputes and silent customers go to a human",
    "Every action is in the audit log: who, what, when"],
      sub="Trust is the product")
slide("Live demo", [
    "Upload invoices - errors flagged",
    "Run agent - reminders drafted",
    "Approve - email sent",
    "Customer replies - status updates / escalation",
    "Weekly report + audit log"], sub="3 minutes")
slide("Built for what matters to Red Rock", [
    "Data intake and validation  ->  invoice extraction + VAT/TRN checks",
    "Task automation  ->  drafts reminders by itself",
    "Deadline monitoring  ->  overdue + due-soon watch",
    "Communications + escalation  ->  reply handling, humans for the hard cases",
    "Automated reports  ->  weekly cash summary",
    "Human approval + audit log  ->  built in"])
slide("Next steps", [
    "Real email sending (Gmail / Outlook) after approval",
    "Connect to accounting software (e.g. Zoho, QuickBooks)",
    "Arabic reminders, WhatsApp channel",
    "Pilot with a few real small businesses"], sub="PayPilot - get paid, stay in control")
prs.save("PayPilot_pitch.pptx")
print("saved PayPilot_pitch.pptx")
